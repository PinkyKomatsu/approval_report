"""過去報告書の取り込みと、見出しごとの候補（Excel がなくても動く）。

- フォルダ内の .xlsx / .xlsm（サブフォルダを含む。~$ で始まる一時ファイルは除く）を openpyxl で読む
- 雛形と同じ方法で見出しを探し、セクションごとの記載行と項目欄の値を取り出す（見出し基準なので行の増減に強い）
- 記載行は 1 行を 1 件の候補にする。「　　⇒」で始まる続きの行は、直前の行とまとめて 1 件にする
- 全報告書を通して同じ文面の候補は 1 件にまとめ、出現回数を数える
"""
from __future__ import annotations

import json
import os
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

from . import template_parser as tp

CONTINUATION = re.compile(r"^[ 　]*⇒")


def data_dir() -> Path:
    """data フォルダ（exe の隣。書き込めなければ %LOCALAPPDATA%）。"""
    base = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent.parent
    d = base / "data"
    try:
        d.mkdir(parents=True, exist_ok=True)
        probe = d / ".write_test"
        probe.write_text("", encoding="utf-8")
        probe.unlink()
        return d
    except OSError:
        d = Path(os.environ.get("LOCALAPPDATA", Path.home())) / "ApprovalReport" / "data"
        d.mkdir(parents=True, exist_ok=True)
        return d


def find_files(folder) -> list[Path]:
    return sorted(p for p in Path(folder).rglob("*")
                  if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm") and not p.name.startswith("~$"))


def phase_of(title: str) -> str:
    """タイトルの括弧内（設計、構築など）。末尾の「工程」は除く。"""
    m = re.search(r"[（(]([^（）()]+)[）)]", title or "")
    if not m:
        return ""
    s = m.group(1).strip()
    if "○" in s:
        return ""
    return s[:-2] if s.endswith("工程") and len(s) > 2 else s


def date_of(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    m = re.search(r"(\d{4})\s*[/\-.年]\s*(\d{1,2})\s*[/\-.月]\s*(\d{1,2})", tp.norm(v))
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
        except ValueError:
            return ""
    return ""


def read_report(path) -> dict:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        name = tp.find_sheet(wb)
        if name is None:
            raise ValueError("対象のシートがありません")
        ws = wb[name]
        f = tp.parse_ws(ws)
        title = tp.cell_text(ws[f["title"]["cell"]].value)
        raw = {it["key"]: ws[it["cell"]].value for it in f["items"]}
        items = {k: tp.cell_text(v).strip() for k, v in raw.items() if k != "date"}
        sections = {s["key"]: tp.read_rows(ws, s) for s in f["sections"] if s["kind"] != "manual" and s["start"]}
        return {"file": str(path), "name": path.name, "title": title, "phase": phase_of(title),
                "date": date_of(raw.get("date")), "customer": items.get("customer", ""), "items": items,
                "sections": sections}
    finally:
        wb.close()


def import_folder(folder) -> dict:
    reports, errors = [], []
    for p in find_files(folder):
        try:
            reports.append(read_report(p))
        except Exception as e:   # 壊れたファイル・別様式のファイルは飛ばす
            errors.append(f"{p.name}：{e}")
    return {"folder": str(folder), "imported": datetime.now().isoformat(timespec="seconds"),
            "reports": reports, "errors": errors}


def save_cases(data: dict, path) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_cases(path) -> dict | None:
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


# ---------------------------------------------------------------------------
# 候補
# ---------------------------------------------------------------------------

@dataclass
class Candidate:
    rows: list[list[str]]                       # 1 件の候補（続きの行を含む。各行は記入する列ごとの値）
    count: int = 0                              # 出現回数
    phases: list[str] = field(default_factory=list)
    dates: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(" ｜ ".join(c for c in row if c) for row in self.rows)


def _is_continuation(row: list[str]) -> bool:
    filled = [c for c in row if c.strip()]
    if not filled:
        return False
    if len(row) > 1 and not row[0].strip():      # 左右の組で右だけの行は、直前の行の続き
        return True
    return all(CONTINUATION.match(c) for c in filled)


def split_candidates(rows: list[list[str]]) -> list[list[list[str]]]:
    out: list[list[list[str]]] = []
    for row in rows:
        row = [c.rstrip() for c in row]
        if not any(c.strip() for c in row):
            continue
        if out and _is_continuation(row):
            out[-1].append(row)
        else:
            out.append([row])
    return out


def build_candidates(reports: list[dict]) -> dict[str, list[Candidate]]:
    """見出しごとの候補。同じ文面は 1 件にまとめる。"""
    out: dict[str, dict[tuple, Candidate]] = {}
    for rep in reports:
        for key, rows in rep.get("sections", {}).items():
            for cand_rows in split_candidates(rows):
                k = tuple(tuple(r) for r in cand_rows)
                c = out.setdefault(key, {}).setdefault(k, Candidate(rows=[list(r) for r in cand_rows]))
                c.count += 1
                if rep.get("phase"):
                    c.phases.append(rep["phase"])
                if rep.get("date"):
                    c.dates.append(rep["date"])
    return {key: list(d.values()) for key, d in out.items()}


def item_values(reports: list[dict], key: str) -> list[str]:
    """項目欄の過去の値（新しい報告書の順、重複なし）。"""
    out = []
    for rep in sorted(reports, key=lambda r: r.get("date") or "", reverse=True):
        v = (rep.get("items") or {}).get(key, "")
        if v and v not in out:
            out.append(v)
    return out


def phases(reports: list[dict]) -> list[str]:
    out = []
    for rep in reports:
        if rep.get("phase") and rep["phase"] not in out:
            out.append(rep["phase"])
    return out
