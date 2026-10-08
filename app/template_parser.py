"""雛形（工程承認会議 報告書）の解析。Excel がなくても動く（openpyxl と、xlsx の中の XML だけを使う）。

- 項目欄：ラベルを探し、その右のセル（結合セルなら左上）を記入先にする
- 文章セクション：B 列の「■」「○」「【」で始まるセルを見出しとし、次の見出しまでを記入範囲にする
- 記入ガイド：P 列より右の吹き出し図形の文を、図形の行位置から各セクションに対応付ける
- 判断に迷うもの（見出しが見つからない・重複している）は status を「要確認」にする
"""
from __future__ import annotations

import json
import posixpath
import re
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

NEEDS_CHECK = "要確認"
HEADING_MARKS = ("■", "○", "【")
HEADING_COL = 2                    # B 列
GUIDE_MIN_COL = 16                 # P 列（これより右の図形を記入ガイドとみなす）

# 項目欄：キー → (ラベルの候補, 既定のセル)
ITEMS = {
    "date": (("作成日",), "M1"),
    "dept": (("部署",), "M2"),
    "name": (("氏名",), "M3"),
    "customer": (("顧客名",), "C5"),
    "project": (("Prj名", "PJ名", "プロジェクト名"), "C7"),
    "work": (("作業内容",), "C9"),
    "overview": (("プロジェクト概要",), "D14"),
}
ITEM_NAMES = {"date": "作成日", "dept": "部署", "name": "氏名", "customer": "顧客名", "project": "Prj名",
              "work": "作業内容", "overview": "プロジェクト概要"}
TITLE_DEFAULT, MEETING_DEFAULT = "B1", "B12"
# 左右 2 列で 1 組の欄（左の見出し, 右の見出し）
PAIRS = (("前回指摘事項", "対策結果"), ("事前レビュー指摘事項", "対策内容"), ("問題点", "対応策"))
# 罫線で囲まれた表（ツールでは記入しない）
MANUAL_WORDS = ("算定トラブル防止施策", "品質指標", "課題件数", "第三者チェック", "○△×")

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


# ---------------------------------------------------------------------------
# 文字列
# ---------------------------------------------------------------------------

def norm(s) -> str:
    """比較用：全角半角をそろえ、空白を除く。"""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(s or "")))


def heading_key(text: str) -> str:
    """見出しの識別名（先頭の ■ ○ を除く）。雛形と過去報告書の見出しの対応付けに使う。"""
    n = norm(text)
    return n.lstrip("■○") or n


def is_heading(text) -> bool:
    return isinstance(text, str) and text[:1] in HEADING_MARKS


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _label_of(text: str) -> str:
    return norm(text).strip("【】[]：:■○□")


# ---------------------------------------------------------------------------
# 図形（吹き出し）
# ---------------------------------------------------------------------------

def _rels(z: zipfile.ZipFile, part: str) -> dict[str, str]:
    """part の関係ファイルを読み、rId → パーツのパス を返す。"""
    d, name = posixpath.split(part)
    path = posixpath.join(d, "_rels", name + ".rels")
    if path not in z.namelist():
        return {}
    out = {}
    for r in ET.fromstring(z.read(path)).iter(f"{{{_PKG}}}Relationship"):
        t = r.get("Target", "")
        if r.get("TargetMode") == "External":
            continue
        out[r.get("Id")] = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(d, t))
    return out


def sheet_shapes(path, sheet_name: str) -> list[tuple[int, int, str]]:
    """シートの図形のうち文字を持つもの [(列, 行, 文)]。列・行は 1 始まり（図形の左上のセル）。"""
    out = []
    try:
        z = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile):
        return out
    with z:
        wb_rels = _rels(z, "xl/workbook.xml")
        part = None
        for s in ET.fromstring(z.read("xl/workbook.xml")).iter(f"{{{_MAIN}}}sheet"):
            if s.get("name") == sheet_name:
                part = wb_rels.get(s.get(f"{{{_REL}}}id"))
        if not part or part not in z.namelist():
            return out
        s_rels = _rels(z, part)
        for d in ET.fromstring(z.read(part)).iter(f"{{{_MAIN}}}drawing"):
            dpart = s_rels.get(d.get(f"{{{_REL}}}id"))
            if not dpart or dpart not in z.namelist():
                continue
            for anchor in ET.fromstring(z.read(dpart)):
                frm = anchor.find(f"{{{_XDR}}}from")
                if frm is None:
                    continue
                col = int(frm.findtext(f"{{{_XDR}}}col", "0")) + 1
                row = int(frm.findtext(f"{{{_XDR}}}row", "0")) + 1
                for sp in anchor.iter(f"{{{_XDR}}}sp"):
                    paras = ["".join(t.text or "" for t in p.iter(f"{{{_A}}}t")) for p in sp.iter(f"{{{_A}}}p")]
                    text = "\n".join(paras).strip()
                    if text:
                        out.append((col, row, text))
    return out


# ---------------------------------------------------------------------------
# シートの解析
# ---------------------------------------------------------------------------

def find_sheet(wb) -> str | None:
    names = list(wb.sheetnames)
    for n in names:
        if "工程承認会議" in n:
            return n
    others = [n for n in names if "版数管理" not in n]
    return others[0] if others else None


def _merged(ws) -> dict[tuple[int, int], tuple[int, int, int, int]]:
    m = {}
    for rng in ws.merged_cells.ranges:
        box = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = box
    return m


def _last_row(ws) -> int:
    area = ws.print_area
    if isinstance(area, (list, tuple)):
        area = area[0] if area else None
    if area:
        m = re.search(r"\$?[A-Z]+\$?(\d+)\s*$", str(area).split(",")[0])
        if m:
            return int(m.group(1))
    return ws.max_row


def _boxed(ws, r: int, c: int) -> bool:
    b = ws.cell(r, c).border
    return all(getattr(b, s) is not None and getattr(b, s).style for s in ("left", "right", "top", "bottom"))


def _is_grid(ws, start: int, end: int) -> bool:
    """罫線で囲まれた表か（四辺に罫線のあるセルが、2 行以上・3 列以上にわたって 6 個以上）。"""
    rows, cols, n = set(), set(), 0
    for r in range(start, end + 1):
        for c in range(HEADING_COL, GUIDE_MIN_COL):
            if _boxed(ws, r, c):
                n += 1
                rows.add(r)
                cols.add(c)
    return n >= 6 and len(rows) >= 2 and len(cols) >= 3


def addr(r: int, c: int) -> str:
    return f"{get_column_letter(c)}{r}"


def parse_ws(ws, shapes: list[tuple[int, int, str]] | None = None) -> dict:
    """1 枚のシートを解析する（雛形・過去報告書の共通処理）。"""
    merged = _merged(ws)
    last = _last_row(ws)
    texts: dict[tuple[int, int], str] = {}
    for row in ws.iter_rows(min_row=1, max_row=max(last, 1)):
        for cell in row:
            if cell.value is not None and isinstance(cell.value, str) and cell.value.strip():
                texts[(cell.row, cell.column)] = cell.value

    def anchor(r, c):
        box = merged.get((r, c))
        return (box[0], box[1]) if box else (r, c)

    used: set[tuple[int, int]] = set()
    warnings: list[str] = []

    # 項目欄
    items = []
    for key, (labels, default) in ITEMS.items():
        wanted = {norm(x) for x in labels}
        found = [rc for rc, t in texts.items() if _label_of(t) in wanted]
        item = {"key": key, "label": ITEM_NAMES[key], "cell": default, "status": "ok", "note": ""}
        if found:
            r, c = sorted(found)[0]
            box = merged.get((r, c), (r, c, r, c))
            item["cell"] = addr(*anchor(r, box[3] + 1))
            used.update(found)
            if len(found) > 1:
                item.update(status=NEEDS_CHECK, note="ラベルが複数あります")
        else:
            item.update(status=NEEDS_CHECK, note="ラベルが見つかりません（既定のセルを仮に使います）")
        items.append(item)

    # 会議種別（□工程承認会議 □プロジェクト報告会）
    meeting = {"cell": MEETING_DEFAULT, "status": NEEDS_CHECK, "note": "会議種別の欄が見つかりません"}
    for rc, t in sorted(texts.items()):
        if ("□" in t or "■" in t) and "工程承認会議" in t and "報告会" in t:
            meeting = {"cell": addr(*rc), "status": "ok", "note": ""}
            used.add(rc)
            break
    # タイトル（上の 3 行で「工程承認会議」「報告会」「○○」を含むセル）
    title = {"cell": TITLE_DEFAULT, "text": cell_text(ws[TITLE_DEFAULT].value), "status": "ok", "note": ""}
    for rc, t in sorted(texts.items()):
        if rc[0] <= 3 and rc not in used and ("工程承認会議" in t or "報告会" in t or "○○" in t):
            title = {"cell": addr(*rc), "text": t, "status": "ok", "note": ""}
            used.add(rc)
            break

    # 見出し
    pair_right = {norm(r): norm(l) for l, r in PAIRS}
    heads = []
    for r in range(1, last + 1):
        t = texts.get((r, HEADING_COL))
        if not is_heading(t) or (r, HEADING_COL) in used:
            continue
        right = None
        for c in range(HEADING_COL + 1, GUIDE_MIN_COL):
            rt = texts.get((r, c))
            if rt and (is_heading(rt) or heading_key(rt) in pair_right):
                right = (c, rt)
                break
        heads.append((r, t, right))

    sections = []
    seen: dict[str, int] = {}
    for i, (r, t, right) in enumerate(heads):
        end = heads[i + 1][0] - 1 if i + 1 < len(heads) else last
        key = heading_key(t)
        seen[key] = seen.get(key, 0) + 1
        if seen[key] > 1:
            key = f"{key}#{seen[key]}"
        kind = "pair" if right else "text"
        if any(norm(w) in norm(t) for w in MANUAL_WORDS) or (end > r and _is_grid(ws, r + 1, end)):
            kind = "manual"
        # 記入する列：記載のある列（結合セルの左上）のうち最も多い列。左右の組は右の見出しの列で分ける
        bounds = [(HEADING_COL, right[0]), (right[0], GUIDE_MIN_COL)] if right else [(HEADING_COL, GUIDE_MIN_COL)]
        cols = []
        for lo, hi in bounds:
            counts: dict[int, int] = {}
            for rr in range(r + 1, end + 1):
                for c in range(lo, hi):
                    if (rr, c) in texts:
                        counts[c] = counts.get(c, 0) + 1
                        break
            cols.append(max(counts, key=lambda c: (counts[c], -c)) if counts else lo)
        samples = [[cell_text(ws.cell(rr, c).value) for c in cols] for rr in range(r + 1, end + 1)]
        sections.append({
            "key": key, "name": t.strip(), "right_name": right[1].strip() if right else "", "row": r,
            "start": r + 1, "end": end, "kind": kind, "cols": [get_column_letter(c) for c in cols],
            "guide": "", "status": "ok", "note": "", "samples": samples})

    # 重複した見出し
    base = {}
    for s in sections:
        base.setdefault(s["key"].split("#")[0], []).append(s)
    for k, ss in base.items():
        if len(ss) > 1:
            for s in ss:
                if s["kind"] != "manual":
                    s.update(status=NEEDS_CHECK, note="同じ見出しが複数あります")
            warnings.append(f"見出し「{k}」が {len(ss)} か所にあります")
    # 見つからない左右の組の見出し
    keys = {s["key"] for s in sections}
    for l, r_ in PAIRS:
        if norm(l) not in keys:
            sections.append({"key": norm(l), "name": f"■{l}", "right_name": f"■{r_}", "row": None, "start": None,
                             "end": None, "kind": "pair", "cols": [], "guide": "", "status": NEEDS_CHECK,
                             "note": "見出しが見つかりません", "samples": []})

    # 記入ガイド（P 列より右の吹き出し）
    items_guide = []
    placed = [s for s in sections if s["row"]]
    for col, row, text in shapes or []:
        if col < GUIDE_MIN_COL:
            continue
        owner = None
        for s in placed:
            if s["row"] <= row:
                owner = s
        if owner is None:
            items_guide.append(text)
        else:
            owner["guide"] = (owner["guide"] + "\n" + text).strip()

    return {"sheet": ws.title, "last_row": last, "title": title, "meeting": meeting, "items": items,
            "items_guide": "\n".join(items_guide), "sections": sections, "warnings": warnings}


def parse_template(path) -> dict:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    name = find_sheet(wb)
    if name is None:
        raise ValueError("対象のシートが見つかりません。")
    fmt = parse_ws(wb[name], sheet_shapes(path, name))
    if "工程承認会議" not in name:
        fmt["warnings"].insert(0, f"名前に「工程承認会議」を含むシートがないため、「{name}」を使います")
    wb.close()
    return fmt


def read_rows(ws, section: dict) -> list[list[str]]:
    """セクションの記入範囲の各行の値（記入する列ごと）。"""
    if not section.get("start") or section["kind"] == "manual":
        return []
    cols = [column_index_from_string(c) for c in section["cols"]]
    return [[cell_text(ws.cell(r, c).value) for c in cols] for r in range(section["start"], section["end"] + 1)]


def set_target(fmt: dict, target: str, row: int, col: int) -> None:
    """［選択中のセルを記入先にする］：項目欄（item:キー）またはセクションの記入先を指定する。"""
    if target.startswith("item:"):
        for it in fmt["items"]:
            if it["key"] == target[5:]:
                it.update(cell=addr(row, col), status="ok", note="手動で指定")
        return
    if target == "meeting":
        fmt["meeting"].update(cell=addr(row, col), status="ok", note="手動で指定")
        return
    heads = sorted(s["row"] for s in fmt["sections"] if s["row"])
    for s in fmt["sections"]:
        if s["key"] != target:
            continue
        nxt = [h for h in heads if h >= row and h != s["row"]]
        end = (nxt[0] - 1) if nxt else fmt["last_row"]
        cols = [get_column_letter(col)] + s["cols"][1:]
        if s["kind"] == "pair" and len(cols) < 2:
            cols.append(get_column_letter(col + 1))
        s.update(start=row, end=max(end, row), cols=cols, status="ok", note="手動で指定")
        if not s["row"]:
            s["row"] = row - 1


def save_format(fmt: dict, path) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(fmt, ensure_ascii=False, indent=2), encoding="utf-8")


def load_format(path) -> dict | None:
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
