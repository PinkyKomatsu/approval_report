"""画面から使う一連の処理（ひな形・フォーマットの登録、過去事例の取込、照合）。GUI には依存しない。
Excel ファイルは読むだけで、変更しない（登録したファイルはアプリのフォルダにコピーして使う）。"""
from __future__ import annotations

import logging
import shutil
from datetime import date, datetime
from pathlib import Path

from excel_io import inspect, profile as profile_mod, reader

from . import paths
from .db import Database
from .matching import Matcher
from .models import EN_SUFFIX, CaseRow, MasterItem
from .normalize import item_key, norm, wbs_level

log = logging.getLogger("approvalreport")


class DuplicateCase(Exception):
    pass


def _copy_in(path: Path, kind: str) -> Path:
    dest_dir = paths.formats_dir() / kind
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"{datetime.now():%Y%m%d%H%M%S%f}_{path.name}"
    shutil.copy2(path, dest)
    return dest


def text_of(v) -> str:
    """セルの値を文字列にする（日付は yyyy/mm/dd）。"""
    if isinstance(v, datetime):
        return f"{v:%Y/%m/%d}" if (v.hour, v.minute, v.second) == (0, 0, 0) else f"{v:%Y/%m/%d %H:%M}"
    if isinstance(v, date):
        return f"{v:%Y/%m/%d}"
    return reader.cell_text(v)


def seed_glossary(db: Database) -> int:
    """用語辞書の初期値（resources/glossary.csv）を入れる。英語が登録済みの語は変えない。"""
    import csv
    current = db.glossary()
    terms = {}
    with open(paths.resource("glossary.csv"), encoding="utf-8-sig", newline="") as f:
        for row in csv.reader(f):
            if len(row) >= 2 and row[0].strip() and row[0].strip() != "日本語" and not current.get(row[0].strip()):
                terms[row[0].strip()] = row[1].strip()
    if terms:
        db.upsert_terms(terms, source="初期値", keep_existing_en=False)
    return len(terms)


# ---------------------------------------------------------------------------
# ひな形・フォーマット
# ---------------------------------------------------------------------------

def build_master_items(rows: list[tuple[int, dict]]) -> list[MasterItem]:
    """ひな形の明細から作業項目のマスタを作る。階層は WBS 番号から（親は直前の浅い階層の項目）。"""
    items: list[MasterItem] = []
    keys: set[str] = set()
    current_phase = ""
    for seq, (r, v) in enumerate(rows):
        name = text_of(v.get("name"))
        if not name:
            continue
        wbs = text_of(v.get("wbs_no"))
        level = wbs_level(wbs) if wbs else (items[-1].level if items else 1)
        phase = text_of(v.get("phase"))
        if level == 1 and not phase:
            current_phase = name
        phase = phase or current_phase
        key = item_key(phase, name)
        if key in keys:
            key = f"{key}#{wbs or seq}"
        keys.add(key)
        parent = next((i.key for i in reversed(items) if i.level < level), None)
        items.append(MasterItem(key=key, seq=seq, row=r, wbs_no=wbs, phase=phase, name=name, norm=norm(name),
                                level=level, parent_key=parent))
    return items


def register_master(db: Database, path, cfg: dict, profile: dict | None = None) -> dict:
    p = reader.check(path)
    stored = _copy_in(p, "master")
    prof = profile or profile_mod.detect(stored, cfg)
    fid = db.add_format("master", p.stem, stored, p, reader.sha256(p), prof, {})
    items = build_master_items(profile_mod.read_rows(stored, prof, cfg.get("profile", {}).get("blank_rows_to_stop", 5)))
    db.replace_master_items(fid, items)
    db.upsert_terms({**{i.name: "" for i in items}, **{i.phase: "" for i in items if i.phase}}, source="ひな形")
    rematch(db, cfg)
    log.info("ひな形を登録しました（%d項目）", len(items))
    return {"format_id": fid, "profile": prof, "items": len(items)}


def reload_master(db: Database, cfg: dict, profile: dict) -> int:
    f = db.active_format("master")
    db.update_format(f["id"], profile)
    items = build_master_items(profile_mod.read_rows(f["file"], profile))
    db.replace_master_items(f["id"], items)
    db.upsert_terms({i.name: "" for i in items}, source="ひな形")
    rematch(db, cfg)
    return len(items)


def register_format(db: Database, path, cfg: dict, profile: dict | None = None) -> dict:
    p = reader.check(path)
    stored = _copy_in(p, "format")
    prof = profile or profile_mod.detect(stored, cfg)
    analysis = inspect.analyze_format(stored, prof) if prof.get("columns") else {}
    fid = db.add_format("format", p.stem, stored, p, reader.sha256(p), prof, analysis)
    log.info("フォーマットを登録しました（%d列）", len(prof.get("columns", {})))
    return {"format_id": fid, "profile": prof, "analysis": analysis}


def update_format_profile(db: Database, profile: dict) -> dict:
    f = db.active_format("format")
    analysis = inspect.analyze_format(f["file"], profile) if profile.get("columns") else {}
    db.update_format(f["id"], profile, analysis)
    return analysis


def set_target(db: Database, item_key_: str, meaning: str | None, ref: str) -> None:
    """記載先のセルを利用者が指定する（プロファイルに保存）。meaning が None なら行だけを指定する。"""
    f = db.active_format("format")
    prof = dict(f["profile"])
    targets = dict(prof.get("targets") or {})
    t = dict(targets.get(item_key_) or {})
    import re
    m = re.match(r"([A-Z]+)(\d+)$", ref)
    if meaning:
        t[meaning] = ref
    if m:
        t["_row"] = int(m.group(2))
    targets[item_key_] = t
    prof["targets"] = targets
    db.update_format(f["id"], prof)


# ---------------------------------------------------------------------------
# 過去事例
# ---------------------------------------------------------------------------

def case_profile(db: Database, path, cfg: dict) -> dict:
    """登録済みのフォーマットのプロファイルで読む（シート・見出しが合わなければ自動推定）。
    フォーマットにない英語の列（見出しに「英語」「(EN)」などを含む列）は、過去事例の見出しから追加する。"""
    f = db.active_format("format")
    prof = dict(f["profile"]) if f else None
    wb = reader.load(path, data_only=True)
    try:
        if prof and prof.get("sheet") in wb.sheetnames and prof.get("header_row"):
            ws = wb[prof["sheet"]]
            heads = prof.get("headers", {})
            same = sum(1 for c, h in heads.items() if norm(reader.cell_text(ws[f"{c}{prof['header_row']}"].value)) == norm(h))
            if heads and same >= max(1, len(heads) * 2 // 3):
                cols = dict(prof["columns"])
                from openpyxl.utils import get_column_letter
                for c in range(1, min(ws.max_column, 60) + 1):
                    letter = get_column_letter(c)
                    meaning, _s = profile_mod.match_header(reader.cell_text(ws.cell(prof["header_row"], c).value))
                    if meaning and meaning.endswith(EN_SUFFIX) and letter not in cols:
                        cols[letter] = meaning
                prof["columns"] = cols
                return prof
    finally:
        wb.close()
    return profile_mod.detect(path, cfg)


def import_case(db: Database, path, cfg: dict, attrs: dict | None = None, name: str | None = None,
                allow_duplicate: bool = False) -> int:
    p = reader.check(path)
    h = reader.sha256(p)
    dup = db.case_by_hash(h)
    if dup and not allow_duplicate:
        raise DuplicateCase(f"「{p.name}」は取込済みです（事例「{dup['name']}」と同じファイル）。")
    prof = case_profile(db, p, cfg)
    rows = []
    phase_now = ""
    for seq, (r, v) in enumerate(profile_mod.read_rows(p, prof, cfg.get("profile", {}).get("blank_rows_to_stop", 5))):
        nm = text_of(v.get("name"))
        if not nm:
            continue
        phase = text_of(v.get("phase")) or phase_now
        phase_now = phase
        texts, texts_en = {}, {}
        for meaning, val in v.items():
            t = text_of(val)
            if not t or meaning in ("name", "wbs_no", "phase", "no"):
                continue
            if meaning.endswith(EN_SUFFIX):
                texts_en[meaning[:-len(EN_SUFFIX)]] = t
            else:
                texts[meaning] = t
        rows.append(CaseRow(id=None, case_id=None, seq=seq, row=r, wbs_no=text_of(v.get("wbs_no")), phase=phase,
                            name=nm, norm=norm(nm), texts=texts, texts_en=texts_en))
    cid = db.add_case(name or p.stem, p, h, attrs or {}, rows)
    rematch(db, cfg, case_id=cid)
    log.info("過去事例を取り込みました（%d行）", len(rows))
    return cid


def rematch(db: Database, cfg: dict, case_id: int | None = None) -> None:
    """照合し直す（照合確認で確定した行はそのまま）。"""
    items = db.master_items()
    matcher = Matcher(items, db.synonyms(), cfg.get("match", {}))
    rows = db.case_rows(case_id)
    results = []
    for m in matcher.match_rows([r for r in rows if r.status != "confirmed" and r.method != "manual"]):
        results.append(m)
    db.save_matches(results)


def confirm_match(db: Database, row_id: int, item_key_: str | None, remember: bool = True) -> None:
    """照合確認での選択。item_key_ が None なら「ひな形外の追加項目」として確定。選んだ結果は同義語辞書に保存する。"""
    row = next(r for r in db.case_rows() if r.id == row_id)
    if item_key_:
        db.set_row_match(row_id, item_key_, "confirmed")
        if remember:
            db.add_synonym(row.norm, item_key_)
    else:
        db.set_row_match(row_id, None, "extra")
