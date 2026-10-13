"""過去事例の列の自動推定（見出しから、対象シート・ヘッダー行・データ開始行・各列の欄を判別する）。

プロファイル（dict）
  sheet        対象シート
  header_row   ヘッダー行
  data_start   データ開始行
  columns      {"B": "wbs_no", "D": "name", "F": "report", "G": "report_en", ...}（列 → 欄）
  headers      {"B": "WBS番号", ...}（列 → 見出しの文字列）
  confidence   0〜1
  issues       迷った点の説明（確信度が閾値未満ならプレビューで利用者が指定する）

見出しに「英語」「(EN)」「English」などを含む列は、同じ欄の英語の列（例：report_en）とみなす。
"""
from __future__ import annotations

import re

from openpyxl.utils import get_column_letter

from core.models import EN_SUFFIX
from core.normalize import norm

from . import reader

# 欄と見出しのキーワード（前にあるほど優先。完全一致は部分一致より強い）
KEYWORDS = {
    "wbs_no": ["wbs番号", "wbsno", "wbs", "作業番号", "タスク番号", "項番"],
    "no": ["no", "№", "#", "連番"],
    "phase": ["工程", "フェーズ", "phase", "大項目", "区分"],
    "name": ["作業項目", "作業項目名", "作業名", "タスク名", "項目名", "タスク", "項目", "task"],
    "work": ["作業内容", "作業概要", "実施内容", "内容"],
    "report": ["報告内容", "報告事項", "報告", "report"],
    "progress": ["進捗状況", "進捗", "状況", "ステータス", "status", "progress"],
    "issue": ["課題懸念", "課題・懸念", "課題", "懸念", "問題点", "リスク", "issue"],
    "policy": ["対応方針", "対応策", "対策", "方針", "対応"],
    "owner": ["担当者", "担当", "責任者", "owner"],
    "plan_date": ["予定日", "完了予定日", "予定", "期限", "plandate"],
    "actual_date": ["実績日", "完了日", "実績", "actualdate"],
    "note": ["備考", "メモ", "コメント", "note", "remarks"],
}
_EN_RE = re.compile(r"[（(]?\s*(英語|英文|英訳|english|en)\s*[)）]?", re.I)
_TOTAL_WORDS = ("合計", "小計", "総計", "total")


def keywords(cfg: dict | None) -> dict[str, list[str]]:
    """設定の header_keywords を、組み込みのキーワードより優先して使う。"""
    extra = {m: [norm(w) for w in ws] for m, ws in (cfg or {}).get("header_keywords", {}).items() if isinstance(ws, list)}
    return {m: extra.get(m, []) + [w for w in KEYWORDS.get(m, []) if w not in extra.get(m, [])]
            for m in list(extra) + [m for m in KEYWORDS if m not in extra]}


def is_english_header(text: str) -> tuple[bool, str]:
    """(英語の列か, 英語を示す語を除いた見出し)"""
    t = str(text or "")
    m = _EN_RE.search(t)
    if not m:
        return False, t
    rest = (t[:m.start()] + t[m.end():]).strip()
    return bool(rest), rest


def match_header(text: str, kw: dict[str, list[str]] | None = None) -> tuple[str | None, float]:
    en, base = is_english_header(text)
    t = norm(base)
    if not t or len(t) > 20:
        return None, 0.0
    best, best_s = None, 0.0
    for meaning, words in (kw or KEYWORDS).items():
        for rank, w in enumerate(words):
            s = 1.0 if t == w else (0.8 if w in t and len(w) >= 2 else 0.0)
            s -= rank * 0.01
            if s > best_s:
                best, best_s = meaning, s
    if best and en:
        return best + EN_SUFFIX, best_s
    return best, best_s


def _header_candidates(ws, scan_rows: int, kw=None):
    cands = []
    for r in range(1, min(ws.max_row, scan_rows) + 1):
        found: dict[str, tuple[str, float, str]] = {}
        dup = []
        for c in range(1, min(ws.max_column, 60) + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            meaning, s = match_header(v, kw)
            if not meaning:
                continue
            col = get_column_letter(c)
            if meaning in found:
                dup.append(meaning)
                if s <= found[meaning][1]:
                    continue
            found[meaning] = (col, s, v.strip())
        score = sum(s for _, s, _ in found.values()) + (1.0 if "name" in found else 0.0)
        if found:
            cands.append({"row": r, "found": found, "score": score, "dup": dup})
    return sorted(cands, key=lambda c: -c["score"])


def _is_blank_row(ws, r: int, cols: list[int]) -> bool:
    return all(ws.cell(r, c).value in (None, "") for c in cols)


def detect(path, cfg: dict | None = None) -> dict:
    cfg = cfg or {}
    scan = int(cfg.get("profile", {}).get("header_scan_rows", 30))
    kw = keywords(cfg)
    wb = reader.load(path, data_only=False)
    ranking = []
    for ws in wb.worksheets:
        for cand in _header_candidates(ws, scan, kw)[:3]:
            ranking.append((cand["score"], ws.title, cand))
    ranking.sort(key=lambda x: -x[0])
    if not ranking:
        return {"sheet": wb.worksheets[0].title, "header_row": None, "data_start": None, "columns": {}, "headers": {},
                "confidence": 0.0, "sheets": wb.sheetnames,
                "issues": ["ヘッダー行が見つかりません。プレビューでヘッダー行と各列を指定してください。"]}
    score, sheet, best = ranking[0]
    ws = wb[sheet]
    found = best["found"]
    columns = {col: m for m, (col, _, _) in found.items()}
    headers = {col: text for m, (col, _, text) in found.items()}
    issues = []
    conf = 0.0
    if "name" in found:
        conf += 0.5
    else:
        issues.append("作業項目の列が見つかりません。")
    if "wbs_no" in found:
        conf += 0.2
    conf += min(3, len([m for m in found if m not in ("name", "wbs_no")])) * 0.1
    if len(ranking) > 1 and ranking[1][0] >= score - 0.3 and (ranking[1][1], ranking[1][2]["row"]) != (sheet, best["row"]):
        conf -= 0.2
        issues.append(f"ヘッダー行の候補が複数あります（{ranking[1][1]} の {ranking[1][2]['row']} 行目）。")
    for m in best["dup"]:
        conf -= 0.1
        issues.append(f"「{m}」に当てはまる列が複数あります。")
    header_row = best["row"]
    data_start = header_row + 1
    col_idx = [ws[f"{c}1"].column for c in columns]
    while data_start < header_row + 4 and _is_blank_row(ws, data_start, col_idx):   # 2 段見出しの隙間などは飛ばす
        data_start += 1
    conf = round(max(0.0, min(1.0, conf)), 2)
    return {"sheet": sheet, "header_row": header_row, "data_start": data_start,
            "columns": dict(sorted(columns.items(), key=lambda kv: (len(kv[0]), kv[0]))),
            "headers": headers, "confidence": conf, "issues": issues, "sheets": wb.sheetnames}



def read_rows(path, profile: dict, blank_stop: int = 5) -> list[tuple[int, dict]]:
    """プロファイルに従って明細を読む。[(行番号, {欄: 値})]（合計行は除く）"""
    if not profile.get("columns") or not profile.get("data_start"):
        return []
    out = []
    for r, values in reader.iter_rows(path, profile["sheet"], profile["data_start"], profile["columns"], blank_stop):
        if any(any(w in norm(reader.cell_text(v)) for w in _TOTAL_WORDS) for v in values.values() if isinstance(v, str)):
            continue
        out.append((r, values))
    return out
