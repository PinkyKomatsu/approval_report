"""フォーマットの解析（記載先のプレビュー・各欄の書式制約・コピー時の列位置に使う）。ファイルは変更しない。

記録するもの（データ開始行のセルを基準にする）
  fields        欄ごとの制約：入力規則の許可値、結合セル、列幅（全角換算の文字数）と行数、日付・数値の表示形式
  formula_cols  数式の列（コピー時に空欄で埋める）
  covered_cols  結合セルの左上以外の列（コピー時に空欄で埋める）
  rows          フォーマットにすでにある作業項目の行（記載先の行を探すのに使う）
"""
from __future__ import annotations

import re
import zipfile

from lxml import etree
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from core.normalize import norm, norm_wbs

from . import reader

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def sheet_part(zf: zipfile.ZipFile, sheet: str) -> str:
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels.iter("{%s}Relationship" % NS_PKG)}
    for s in wb.iter("{%s}sheet" % NS_MAIN):
        if s.get("name") == sheet:
            t = targets[s.get("{%s}id" % NS_R)]
            return t.lstrip("/") if t.startswith("/") else "xl/" + t
    raise KeyError(f"シート「{sheet}」が見つかりません")


def _allowed_values(wb, dv) -> list[str] | None:
    """リストの入力規則の許可値。求められなければ None。"""
    f = (dv.formula1 or "").strip()
    if not f:
        return None
    if f.startswith('"') and f.endswith('"'):
        return [x.strip() for x in f[1:-1].split(",")]
    ref = f.lstrip("=")
    if ref in wb.defined_names:
        ref = wb.defined_names[ref].attr_text
    m = re.match(r"^'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", ref)
    if not m:
        return None
    sheet = m.group(1)
    if sheet not in wb.sheetnames:
        return None
    rng = f"{m.group(2)}{m.group(3)}" + (f":{m.group(4)}{m.group(5)}" if m.group(4) else "")
    cells = wb[sheet][rng]
    if not isinstance(cells, tuple):
        cells = ((cells,),)
    vals = []
    for row in cells:
        for c in (row if isinstance(row, tuple) else (row,)):
            if c.value not in (None, ""):
                vals.append(reader.cell_text(c.value))
    return vals


NS_X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
NS_XM = "http://schemas.microsoft.com/office/excel/2006/main"


def _ext_validations(zf: zipfile.ZipFile, part: str) -> list[dict]:
    """拡張領域（x14:dataValidation）の入力規則。別シートを参照するリストは Excel がここに保存し、openpyxl は読まない。"""
    root = etree.fromstring(zf.read(part))
    out = []
    for dv in root.iter("{%s}dataValidation" % NS_X14):
        f1 = dv.find("{%s}formula1/{%s}f" % (NS_X14, NS_XM))
        sq = dv.find("{%s}sqref" % NS_XM)
        out.append({"type": dv.get("type"), "formula1": f1.text if f1 is not None else None,
                    "sqref": sq.text if sq is not None else ""})
    return out


class _DV:   # _allowed_values に渡すための入れ物
    def __init__(self, formula1):
        self.formula1 = formula1


def _validations(wb, ws, path, sheet) -> list[dict]:
    out = []
    for dv in ws.data_validations.dataValidation:
        out.append({"sqref": str(dv.sqref), "type": dv.type,
                    "allowed": _allowed_values(wb, dv) if dv.type == "list" else None})
    with zipfile.ZipFile(path) as zf:
        for x in _ext_validations(zf, sheet_part(zf, sheet)):
            out.append({"sqref": x["sqref"], "type": x["type"],
                        "allowed": _allowed_values(wb, _DV(x["formula1"])) if x["type"] == "list" else None})
    return out


def _in(ref_ranges: str, col: int, row: int) -> bool:
    for part in (ref_ranges or "").split():
        min_c, min_r, max_c, max_r = range_boundaries(part)
        if min_c <= col <= max_c and min_r <= row <= max_r:
            return True
    return False


def analyze_format(path, profile: dict) -> dict:
    p = reader.check(path)
    wb = reader.load(p, data_only=False)
    sheet = profile["sheet"]
    ws = wb[sheet]
    r0 = profile.get("data_start") or ((profile.get("header_row") or 0) + 1)
    merged = [(m.min_col, m.min_row, m.max_col, m.max_row) for m in ws.merged_cells.ranges]
    validations = _validations(wb, ws, p, sheet)
    cols = {column_index_from_string(c): m for c, m in profile.get("columns", {}).items()}
    max_col = max(list(cols) + [1])

    def merge_of(c, r):
        return next((m for m in merged if m[0] <= c <= m[2] and m[1] <= r <= m[3]), None)

    formula_cols, covered_cols = [], []
    for c in range(1, max_col + 1):
        cell = ws.cell(r0, c)
        if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
            formula_cols.append(get_column_letter(c))
        m = merge_of(c, r0)
        if m and (m[0], m[1]) != (c, r0):
            covered_cols.append(get_column_letter(c))
    fields = {}
    for c, meaning in cols.items():
        m = merge_of(c, r0)
        span = range(m[0], m[2] + 1) if m else [c]
        width = 0.0
        for cc in span:
            dim = ws.column_dimensions.get(get_column_letter(cc))
            width += dim.width if dim is not None and dim.width else 8.43
        height = ws.row_dimensions[r0].height if r0 in ws.row_dimensions and ws.row_dimensions[r0].height else 15.0
        rows_in_merge = (m[3] - m[1] + 1) if m else 1
        v = next((x for x in validations if _in(x["sqref"], c, r0)), None)
        cell = ws.cell(r0, c)
        fields[meaning] = {
            "col": get_column_letter(c),
            "allowed": (v or {}).get("allowed"),
            "number_format": cell.number_format if cell.number_format and cell.number_format != "General" else "",
            "width_chars": int(width // 2),                       # 全角 1 字 ≒ 列幅 2
            "lines": max(1, round(height * rows_in_merge / 15.0)),
            "merged": (f"{get_column_letter(m[0])}{m[1]}:{get_column_letter(m[2])}{m[3]}" if m else ""),
            "wrap": bool(cell.alignment is not None and cell.alignment.wrap_text),
        }
    # フォーマットにすでにある作業項目の行
    rows = []
    name_col = next((c for c, m in cols.items() if m == "name"), None)
    wbs_col = next((c for c, m in cols.items() if m == "wbs_no"), None)
    if name_col or wbs_col:
        for r in range(r0, ws.max_row + 1):
            name = reader.cell_text(ws.cell(r, name_col).value) if name_col else ""
            wbs = reader.cell_text(ws.cell(r, wbs_col).value) if wbs_col else ""
            if name or wbs:
                rows.append({"row": r, "name": norm(name), "wbs_no": norm_wbs(wbs)})
    wb.close()
    return {"sheet": sheet, "data_start": r0, "max_col": get_column_letter(max_col), "formula_cols": formula_cols,
            "covered_cols": covered_cols, "fields": fields, "rows": rows, "validations": validations,
            "merged": [f"{get_column_letter(a)}{b}:{get_column_letter(c)}{d}" for a, b, c, d in merged]}


def target_cells(analysis: dict, profile: dict, item_key: str, wbs_no: str, name: str, meaning: str) -> tuple[list[str], str]:
    """記載先のセル候補と状態（"ok" / "picked"＝利用者が指定 / "ambiguous"＝複数 / "missing"＝見つからない）。"""
    picked = (profile.get("targets") or {}).get(item_key, {})
    col = analysis.get("fields", {}).get(meaning, {}).get("col")
    if picked.get(meaning):
        return [picked[meaning]], "picked"
    if picked.get("_row") and col:
        return [f"{col}{picked['_row']}"], "picked"
    if not col:
        return [], "missing"
    w, n = norm_wbs(wbs_no), norm(name)
    hits = [r["row"] for r in analysis.get("rows", []) if (w and r["wbs_no"] == w)]
    if len(hits) != 1:
        by_name = [r["row"] for r in analysis.get("rows", []) if n and r["name"] == n]
        hits = by_name if (len(by_name) == 1 or not hits) else hits
    if len(hits) == 1:
        return [f"{col}{hits[0]}"], "ok"
    return [f"{col}{h}" for h in hits], ("ambiguous" if hits else "missing")
