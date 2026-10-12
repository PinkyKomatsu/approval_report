"""プレビュー用のデータ（セルの値・塗り・太字・結合・列幅）。表示専用で、ファイルは変更しない。"""
from __future__ import annotations

from openpyxl.utils import get_column_letter

from . import reader


def sheet_grid(path, sheet: str, max_rows: int = 300, max_cols: int = 30) -> dict:
    wb = reader.load(path, data_only=True)
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
        n_rows, n_cols = min(ws.max_row, max_rows), min(ws.max_column, max_cols)
        cells = {}
        for r in range(1, n_rows + 1):
            for c in range(1, n_cols + 1):
                cell = ws.cell(r, c)
                fill = None
                try:
                    if cell.fill is not None and cell.fill.fill_type == "solid" and isinstance(cell.fill.fgColor.rgb, str):
                        fill = "#" + cell.fill.fgColor.rgb[-6:]
                except (AttributeError, TypeError):
                    pass
                text = reader.cell_text(cell.value)
                bold = bool(cell.font is not None and cell.font.b)
                if text or fill or bold:
                    cells[(r, c)] = {"text": text, "fill": fill, "bold": bold}
        spans = [(m.min_row, m.min_col, min(m.max_row, n_rows) - m.min_row + 1, min(m.max_col, n_cols) - m.min_col + 1)
                 for m in ws.merged_cells.ranges if m.min_row <= n_rows and m.min_col <= n_cols]
        widths = {}
        for c in range(1, n_cols + 1):
            dim = ws.column_dimensions.get(get_column_letter(c))
            widths[c] = dim.width if dim is not None and dim.width else 8.43
        return {"rows": n_rows, "cols": n_cols, "cells": cells, "spans": spans, "widths": widths, "sheet": ws.title}
    finally:
        wb.close()
