"""テスト用：雛形 V1.14 の構造をまねた架空の xlsx を作る（実在の顧客名は使わない）。

- 項目欄：B1 タイトル、L1〜L3 ラベル → M1〜M3、B5/B7/B9 ラベル → C5/C7/C9、B12 会議種別、B14:C14 ラベル → D14
- 文章セクション（B 列の ■ ○ 【 で始まる見出し）、左右 2 列の組、罫線で囲まれた表（数式あり）
- コメント、結合セル、印刷範囲、P 列より右の吹き出し図形（記入ガイド）
"""
from __future__ import annotations

import re
import sys
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")

# (見出し, 種類, 右の見出し, 記入例の行（雛形）, 最低の行数)
LAYOUT = [
    ("■作業体制", "text", None, [["　・PM：○○、PL：○○"], ["　・医事担当：○○"]], 2),
    ("■進捗状況", "text", None, [], 0),
    ("○進捗【医事】", "text", None, [["　・WBS進捗：XX/XX 完了"], ["　　⇒遅延項目：○○"]], 2),
    ("○進捗【DPC】", "text", None, [["　・WBS進捗：XX/XX 完了"]], 1),
    ("■前回指摘事項", "pair", "■対策結果", [["　・○○について", "　・○○を実施"]], 2),
    ("■事前レビュー指摘事項", "pair", "■対策内容", [["　・○○", "　・○○"]], 2),
    ("■算定トラブル防止施策", "table", None, None, 3),
    ("■品質指標", "table", None, None, 3),
    ("■問題点", "pair", "■対応策", [["　・○○", "　・○○"]], 2),
    ("■今後の予定", "text", None, [["　・○月○日 ○○"], ["　・○月○日 ○○"]], 2),
]

GUIDES = {   # 見出し → 吹き出しの文
    "■作業体制": "体制図の役割と氏名を記載する",
    "○進捗【医事】": "WBS の完了数／全体数と、遅延している項目を記載する",
    "■問題点": "未解決の問題点と対応策を左右に記載する",
    "■今後の予定": "次回会議までの主要な予定を日付付きで記載する",
}
ITEMS_GUIDE = "部署・氏名は報告者を記載する"


def build_report(path, *, title="20XX年XX月 工程承認会議（○○工程）", made=None, customer="", project="",
                 work="", dept="", name="", overview="", lines: dict | None = None, extra_rows: int = 0,
                 shapes: bool = True, duplicate: str | None = None, drop: tuple = ()) -> dict:
    """報告書（雛形・過去報告書）を作り、見出しの行番号などを返す。"""
    lines = lines or {}
    wb = openpyxl.Workbook()
    vs = wb.active
    vs.title = "版数管理"
    vs["A1"] = "版数"
    vs["B1"] = "V1.14"
    ws = wb.create_sheet("工程承認会議")
    for col, w in zip("ABCDEFGHIJKLMNO", (2, 14, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 12, 12)):
        ws.column_dimensions[col].width = w
    ws["B1"] = title
    ws["B1"].font = Font(size=14, bold=True)
    ws.merge_cells("B1:J1")
    for r, label, v in ((1, "作成日", made), (2, "部署", dept), (3, "氏名", name)):
        ws.cell(r, 12, label).border = BOX
        ws.cell(r, 13, v or None).border = BOX
        ws.merge_cells(start_row=r, start_column=13, end_row=r, end_column=15)
    ws["M1"].number_format = "yyyy/mm/dd"
    ws["M1"].comment = Comment("会議の開催日ではなく作成日を記入", "作成者")
    for r, label, v in ((5, "顧客名", customer), (7, "Prj名", project), (9, "作業内容", work)):
        ws.cell(r, 2, label).fill = HEAD_FILL
        ws.cell(r, 3, v or None)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=15)
    ws["B12"] = "□工程承認会議　□プロジェクト報告会"
    ws["B14"] = "プロジェクト概要"
    ws.merge_cells("B14:C14")
    ws["D14"] = overview or None
    ws.merge_cells("D14:O14")

    rows: dict[str, int] = {}
    r = 16
    layout = [x for x in LAYOUT if x[0] not in drop]
    if duplicate:
        layout = layout + [x for x in LAYOUT if x[0] == duplicate]
    for head, kind, right, sample, base in layout:
        rows.setdefault(head, r)
        ws.cell(r, 2, head).font = Font(bold=True)
        if kind == "pair":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
            ws.cell(r, 8, right).font = Font(bold=True)
            ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=15)
        else:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=15)
        r += 1
        if kind == "table":
            heads = ["No", "施策・指標", "", "", "計画", "実績", "", "", "", "", "", "", "合計", "備考"]
            for i in range(base):
                for c in range(2, 16):
                    cell = ws.cell(r + i, c)
                    cell.border = BOX
                    if i == 0:
                        cell.value = heads[c - 2] or None
                    elif c == 2:
                        cell.value = i
                    elif c == 14:
                        cell.value = f"=F{r + i}+G{r + i}"
            r += base
            continue
        body = lines.get(head, sample)
        n = max(len(body), base) + (extra_rows if base else 0)
        for i in range(n):
            vals = body[i] if i < len(body) else [""] * (2 if kind == "pair" else 1)
            if kind == "pair":
                ws.cell(r + i, 2, vals[0] or None)
                ws.merge_cells(start_row=r + i, start_column=2, end_row=r + i, end_column=7)
                ws.cell(r + i, 8, vals[1] or None)
                ws.merge_cells(start_row=r + i, start_column=8, end_row=r + i, end_column=15)
            else:
                ws.cell(r + i, 2, vals[0] or None).alignment = Alignment(vertical="top")
                ws.merge_cells(start_row=r + i, start_column=2, end_row=r + i, end_column=15)
            ws.cell(r + i, 2).border = Border(left=THIN)
            ws.cell(r + i, 15).border = Border(right=THIN)
        r += n
    last = r - 1
    ws.print_area = f"A1:O{last}"
    wb.save(path)
    if shapes:
        anchors = [(16, 2, ITEMS_GUIDE)] + [(16, rows[h], t) for h, t in GUIDES.items() if h in rows]
        anchors.append((3, 5, "P 列より左の図形（ガイドではない）"))
        add_shapes(path, "xl/worksheets/sheet2.xml", anchors)
    return {"rows": rows, "last": last}


def _shape(i: int, col: int, row: int, text: str) -> str:
    return (f'<xdr:twoCellAnchor editAs="oneCell"><xdr:from><xdr:col>{col - 1}</xdr:col><xdr:colOff>0</xdr:colOff>'
            f'<xdr:row>{row - 1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>{col + 3}</xdr:col>'
            f'<xdr:colOff>0</xdr:colOff><xdr:row>{row + 1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
            f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{i + 1}" name="吹き出し {i}"/><xdr:cNvSpPr/>'
            f'</xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm>'
            f'<a:prstGeom prst="wedgeRectCallout"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="FFF2CC"/>'
            f'</a:solidFill><a:ln><a:solidFill><a:srgbClr val="BF9000"/></a:solidFill></a:ln></xdr:spPr>'
            f'<xdr:txBody><a:bodyPr vertOverflow="clip" wrap="square" rtlCol="0" anchor="t"/><a:lstStyle/>'
            f'<a:p><a:r><a:rPr lang="ja-JP" altLang="en-US" sz="900"/><a:t>{text}</a:t></a:r></a:p></xdr:txBody>'
            f'</xdr:sp><xdr:clientData/></xdr:twoCellAnchor>')


def add_shapes(path, sheet_part: str, anchors) -> None:
    """openpyxl では図形を作れないので、xlsx の中に吹き出しの図形（drawing）を追加する。"""
    path = Path(path)
    with zipfile.ZipFile(path) as z:
        parts = {i: z.read(i.filename) for i in z.infolist()}
    by_name = {i.filename: i for i in parts}
    drawing = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
               'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
               + "".join(_shape(i + 1, c, r, t) for i, (c, r, t) in enumerate(anchors)) + "</xdr:wsDr>")
    sheet = parts[by_name[sheet_part]].decode("utf-8")
    tag = '<drawing r:id="rIdShapes1"/>'
    m = re.search(r"<legacyDrawing|<tableParts|<extLst|</worksheet>", sheet)
    sheet = sheet[:m.start()] + tag + sheet[m.start():]
    if "xmlns:r=" not in sheet[:sheet.index(">", sheet.index("<worksheet"))]:
        sheet = sheet.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    rels_name = sheet_part.replace("worksheets/", "worksheets/_rels/") + ".rels"
    rel = ('<Relationship Id="rIdShapes1" Target="../drawings/drawing_shapes1.xml" '
           'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"/>')
    if rels_name in by_name:
        rels = parts[by_name[rels_name]].decode("utf-8").replace("</Relationships>", rel + "</Relationships>")
    else:
        rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships '
                'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + rel + "</Relationships>")
    ct = parts[by_name["[Content_Types].xml"]].decode("utf-8").replace(
        "</Types>", '<Override PartName="/xl/drawings/drawing_shapes1.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in parts.items():
            if info.filename == sheet_part:
                data = sheet.encode("utf-8")
            elif info.filename == rels_name:
                continue
            elif info.filename == "[Content_Types].xml":
                data = ct.encode("utf-8")
            z.writestr(info.filename, data)
        z.writestr(rels_name, rels.encode("utf-8"))
        z.writestr("xl/drawings/drawing_shapes1.xml", drawing.encode("utf-8"))
    tmp.replace(path)


# ---------------------------------------------------------------------------
# 過去報告書 3 件（行数がそれぞれ違う）
# ---------------------------------------------------------------------------

REPORTS = {
    "2025/設計_0510.xlsx": dict(
        title="2025年5月 工程承認会議（設計工程）", made=date(2025, 5, 10), customer="架空総合病院", project="医事会計更新",
        work="基本設計", dept="医療システム部", name="山田 太郎", overview="医事会計システムの更新",
        lines={"■作業体制": [["　・PM：山田、PL：佐藤"], ["　・医事担当：鈴木"]],
               "○進捗【医事】": [["　・WBS進捗：20/40 完了"], ["　・自動精算機メーカとのIF仕様が未確定"],
                               ["　　⇒来月中に仕様確定予定"]],
               "■問題点": [["　・自動精算機IFの仕様未確定", "　・メーカと週次で協議"]],
               "■今後の予定": [["　・6月10日 詳細設計レビュー"], ["　・定例会は毎週実施"]]}),
    "2025/sub/構築_0901.xlsx": dict(
        title="2025年9月 工程承認会議（構築工程）", made=date(2025, 9, 1), customer="架空総合病院", project="医事会計更新",
        work="構築", dept="医療システム部", name="山田 太郎", overview="医事会計システムの更新", extra_rows=2,
        lines={"■作業体制": [["　・PM：山田、PL：佐藤"], ["　・医事担当：鈴木"], ["　・構築担当：高橋"]],
               "○進捗【医事】": [["　・WBS進捗：30/50 完了"], ["　・自動精算機メーカとのIF仕様が未合意"],
                               ["　　⇒9月末の合意を目指す"], ["　・帳票の作成は計画どおり"]],
               "■前回指摘事項": [["　・IF仕様の確定時期", "　・メーカと9月末で合意"], ["", "　　⇒議事録で確認"]],
               "■今後の予定": [["　・定例会は毎週実施"], ["　・10月15日 結合テスト開始"]]}),
    "2026/構築_0115.xlsx": dict(
        title="2026年1月 工程承認会議（構築工程）", made=date(2026, 1, 15), customer="架空市民病院", project="医事会計更新",
        work="構築", dept="医療システム部", name="佐藤 花子", overview="医事会計システムの更新",
        lines={"■作業体制": [["　・PM：山田、PL：佐藤"]],
               "○進捗【医事】": [["　・WBS進捗：45/50 完了"]],
               "■今後の予定": [["　・定例会は毎週実施"]]}),
}


@pytest.fixture
def template(tmp_path) -> Path:
    p = tmp_path / "01_PJ報告書_V1.14.xlsx"
    build_report(p)
    return p


@pytest.fixture
def reports_dir(tmp_path) -> Path:
    d = tmp_path / "過去報告書"
    for rel, kw in REPORTS.items():
        p = d / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        build_report(p, shapes=False, **kw)
    (d / "2026" / "~$構築_0115.xlsx").write_bytes(b"lock")      # Excel の一時ファイル（読まない）
    (d / "メモ.txt").write_text("対象外", encoding="utf-8")
    return d
