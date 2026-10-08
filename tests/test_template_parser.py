import openpyxl

from app import template_parser as tp
from conftest import BOX, ITEMS_GUIDE, build_report


def _by_key(fmt):
    return {s["key"]: s for s in fmt["sections"]}


def test_items_title_meeting(template):
    fmt = tp.parse_template(template)
    assert fmt["sheet"] == "工程承認会議", "版数管理シートは対象外"
    cells = {it["key"]: it["cell"] for it in fmt["items"]}
    assert cells == {"date": "M1", "dept": "M2", "name": "M3", "customer": "C5", "project": "C7", "work": "C9",
                     "overview": "D14"}
    assert all(it["status"] == "ok" for it in fmt["items"])
    assert fmt["title"]["cell"] == "B1" and fmt["meeting"]["cell"] == "B12"
    assert fmt["meeting"]["status"] == "ok" and not fmt["warnings"]


def test_sections_pairs_and_manual_tables(template):
    info = build_report(template)
    fmt = tp.parse_template(template)
    s = _by_key(fmt)
    assert list(s) == ["作業体制", "進捗状況", "進捗【医事】", "進捗【DPC】", "前回指摘事項", "事前レビュー指摘事項",
                       "算定トラブル防止施策", "品質指標", "問題点", "今後の予定"]
    assert s["作業体制"]["kind"] == "text" and s["作業体制"]["cols"] == ["B"]
    r = info["rows"]["■作業体制"]
    assert (s["作業体制"]["start"], s["作業体制"]["end"]) == (r + 1, r + 2)
    assert s["進捗状況"]["start"] > s["進捗状況"]["end"], "直後に ○ 見出しが続く見出しは記入行 0"
    for k, right in (("前回指摘事項", "■対策結果"), ("事前レビュー指摘事項", "■対策内容"), ("問題点", "■対応策")):
        assert s[k]["kind"] == "pair" and s[k]["cols"] == ["B", "H"] and s[k]["right_name"] == right
    assert s["算定トラブル防止施策"]["kind"] == "manual" and s["品質指標"]["kind"] == "manual"
    assert s["今後の予定"]["end"] == info["last"], "最後のセクションは印刷範囲の最終行まで"
    assert s["進捗【医事】"]["samples"] == [["　・WBS進捗：XX/XX 完了"], ["　　⇒遅延項目：○○"]]
    assert all(x["status"] == "ok" for x in fmt["sections"])


def test_guides_are_mapped_by_row(template):
    fmt = tp.parse_template(template)
    s = _by_key(fmt)
    assert s["作業体制"]["guide"] == "体制図の役割と氏名を記載する"
    assert s["進捗【医事】"]["guide"].startswith("WBS の完了数")
    assert s["問題点"]["guide"] == "未解決の問題点と対応策を左右に記載する"
    assert s["今後の予定"]["guide"].startswith("次回会議までの")
    assert s["進捗状況"]["guide"] == "" and s["前回指摘事項"]["guide"] == ""
    assert fmt["items_guide"] == ITEMS_GUIDE, "P 列より左の図形はガイドにしない"


def test_table_detected_by_borders_even_with_unknown_name(tmp_path):
    p = tmp_path / "t.xlsx"
    build_report(p, shapes=False)
    wb = openpyxl.load_workbook(p)
    ws = wb["工程承認会議"]
    s = _by_key(tp.parse_ws(ws))["今後の予定"]
    ws.cell(s["row"], 2).value = "■その他の集計"
    for r in range(s["start"], s["start"] + 2):
        for c in range(2, 6):
            ws.cell(r, c).border = BOX
    sec = {x["name"]: x for x in tp.parse_ws(ws)["sections"]}["■その他の集計"]
    assert sec["kind"] == "manual"


def test_missing_and_duplicate_headings_need_check(tmp_path):
    p = tmp_path / "t.xlsx"
    build_report(p, duplicate="■今後の予定", drop=("■事前レビュー指摘事項",))
    wb = openpyxl.load_workbook(p)
    wb["工程承認会議"]["L2"] = "所属"          # 部署のラベルがない
    wb.save(p)
    fmt = tp.parse_template(p)
    s = _by_key(fmt)
    assert s["今後の予定"]["status"] == tp.NEEDS_CHECK and s["今後の予定#2"]["status"] == tp.NEEDS_CHECK
    assert s["事前レビュー指摘事項"]["status"] == tp.NEEDS_CHECK and s["事前レビュー指摘事項"]["start"] is None
    dept = next(it for it in fmt["items"] if it["key"] == "dept")
    assert dept["status"] == tp.NEEDS_CHECK and dept["cell"] == "M2"
    # ［選択中のセルを記入先にする］
    tp.set_target(fmt, "item:dept", 2, 14)
    assert dept["cell"] == "N2" and dept["status"] == "ok"
    first = s["今後の予定"]
    tp.set_target(fmt, "今後の予定", first["start"], 3)
    assert first["status"] == "ok" and first["cols"] == ["C"] and first["end"] == s["今後の予定#2"]["row"] - 1
    tp.set_target(fmt, "事前レビュー指摘事項", 30, 2)
    rv = s["事前レビュー指摘事項"]
    assert rv["start"] == 30 and rv["cols"] == ["B", "C"] and rv["status"] == "ok"


def test_save_and_load_format(template, tmp_path):
    fmt = tp.parse_template(template)
    tp.save_format(fmt, tmp_path / "data" / "format.json")
    assert tp.load_format(tmp_path / "data" / "format.json") == fmt
